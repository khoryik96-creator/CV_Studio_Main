"""Per-installation Apollo OAuth connection; no Flask or application imports.

The shared broker owns Apollo tokens/client secrets. This installation keeps
its opaque connection credential and revealed emails in the protected vault.
Paid requests are never retried; persisted operation IDs allow result checks.
"""
from __future__ import annotations

import hashlib
import json
import re
import secrets
import threading
import time
from datetime import datetime, timezone
from io import BytesIO
from urllib import request, error
from urllib.parse import urlsplit


USER_AGENT = "CVStudio-Apollo/1.0"


class ApolloError(Exception):
    def __init__(self, message, code="APOLLO_INPUT", status=400):
        super().__init__(message)
        self.code, self.status = code, status


def broker_origin(value):
    value = str(value or "").strip().rstrip("/")
    # A deployment URL is public configuration, never a user-supplied target.
    if value and not re.fullmatch(r"https://[a-z0-9][a-z0-9-]{0,62}\.[a-z0-9][a-z0-9-]{0,62}\.workers\.dev", value):
        raise ApolloError("The Apollo connection address needs owner setup.", "APOLLO_SETUP", 503)
    return value


class _NoRedirect(request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        try:
            fp.close()
        except Exception:  # cleanup-only: preserve the redirect rejection even if close fails
            pass
        raise ApolloError("The Apollo service redirected unexpectedly. Check its setup.", "APOLLO_NETWORK", 502)


class ApolloBrokerClient:
    def __init__(self, origin, opener=None):
        self.origin = broker_origin(origin)
        self.opener = opener or request.build_opener(_NoRedirect()).open

    def post(self, action, body):
        if not self.origin:
            raise ApolloError("Apollo sign-in is not enabled yet. Your CV Studio owner needs to finish setup.", "APOLLO_SETUP", 503)
        if action not in {"start", "poll", "cancel", "disconnect", "reveal", "lookup-status"}:
            raise ApolloError("Choose a supported Apollo action.")
        # Cloudflare refuses urllib's default "Python-urllib" agent (error 1010).
        req = request.Request(self.origin + "/" + action, data=json.dumps(body).encode(),
                              headers={"Content-Type":"application/json", "Accept":"application/json",
                                       "User-Agent":USER_AGENT}, method="POST")
        try:
            with self.opener(req, timeout=25) as response:
                raw = response.read(131073)
                if len(raw) > 131072:
                    raise ValueError()
                data = json.loads(raw)
                if not isinstance(data, dict):
                    raise ValueError()
                return data
        except error.HTTPError as exc:
            status = exc.code
            exc.close()
            messages = {401:"Reconnect your Apollo account.", 403:"Your Apollo account cannot use this integration. Check its permissions.",
                        409:"Apollo is already processing this action. Check the previous lookup before trying again.",
                        429:"Apollo or the sign-in service reached a limit. Try again later."}
            raise ApolloError(messages.get(status, "The Apollo service could not complete this action."), "APOLLO_AUTH" if status==401 else "APOLLO_SERVICE", 502) from None
        except ApolloError:
            raise
        except Exception:
            # Never include provider replies, credentials or profile data in errors.
            raise ApolloError("The Apollo service did not confirm a result. No request was repeated.", "APOLLO_NETWORK", 502) from None


RESULT_MESSAGES = {
    "found":"Business email returned by Apollo. Its status is Apollo's assessment.",
    "not_found":"Apollo did not return a business email for this person.",
    "mismatch":"Apollo could not confirm the same profile and company. No email was attached.",
    "pending":"This lookup has not been confirmed. Check its result before making another paid request.",
    "unconfirmed":"Apollo did not confirm this lookup. Check Apollo before starting a new lookup.",
    "auth_required":"Reconnect Apollo, then check the previous lookup. It will not be repeated automatically.",
    "limited":"Apollo declined the lookup or reached an account limit. Check your Apollo account.",
}


class ApolloService:
    def __init__(self, *, client, load, save, clock=time.time, contacts=None, rows=None):
        self.client, self.load, self.save, self.clock = client, load, save, clock
        self.contacts, self.rows = contacts, rows
        self.lock = threading.Lock()

    def _write(self, value):
        try:
            self.save(value)
        except Exception:
            raise ApolloError("The Apollo connection could not be saved securely. No new lookup will start.", "APOLLO_STORAGE", 503) from None

    def status(self):
        data = self._data()
        return {"available":bool(self.client.origin), "connected":bool(data.get("connection")),
                "name":str(data.get("name") or "")[:160], "busy":self.lock.locked(),
                "connecting":bool(data.get("flow") and data["flow"].get("expires", 0) > self.clock())}

    def _data(self):
        data = self.load()
        # Never forward a credential to a different configured service.
        return data if data.get("broker_url") == self.client.origin else {}

    @staticmethod
    def _credentials(data):
        if not data.get("connection") or not data.get("secret"):
            raise ApolloError("Connect your Apollo account first.", "APOLLO_CONNECT", 409)
        return {"connection_id":data["connection"], "secret":data["secret"]}

    @staticmethod
    def _identity(row, person):
        return hashlib.sha256(json.dumps([row["id"], row["company"], row["hr"]["company_id"], person["url"].rstrip("/")], separators=(",", ":")).encode()).hexdigest()

    def _archive(self, rows, cache):
        if self.contacts:
            self.contacts.remember([(row, person, cache[self._identity(row, person)])
                for row in rows if row.get("hr")
                for person in row["hr"]["people"] if self._identity(row, person) in cache])

    def perform(self, action, *, row=None, url="", confirmed=False, query=""):
        if not self.lock.acquire(blocking=False):
            raise ApolloError("Finish the current Apollo action first.", "APOLLO_BUSY", 409)
        try:
            data = self._data()
            if action == "contacts":
                if not self.contacts:
                    return {"contacts":[]}
                # Read-only recovery of old local results, never a broker lookup.
                cache = data.get("emails", {})
                warning = ""
                if cache and self.rows:
                    try:
                        self._archive(self.rows(), cache)
                    except Exception:
                        # Existing history must still be viewable/exportable if
                        # importing another email hits capacity/ad-store failure.
                        warning = "Some HR-card emails could not be copied to contacts. Saved contacts are shown; check storage and refresh to recover the remaining emails without another credit."
                return {"contacts":self.contacts.list(query), "contact_warning":warning}
            if action == "connect":
                if data.get("connection"):
                    raise ApolloError("Disconnect the current Apollo account before changing accounts.")
                if data.get("flow", {}).get("expires", 0) > self.clock():
                    return {"authorization_url":data["flow"]["url"], "connecting":True}
                secret = secrets.token_hex(32)
                result = self.client.post("start", {"proof_hash":hashlib.sha256(secret.encode()).hexdigest()})
                auth = result.get("authorization_url", "")
                parsed = urlsplit(auth) if isinstance(auth, str) else None
                if (not parsed or parsed.scheme != "https" or parsed.netloc != "app.apollo.io" or
                        parsed.path != "/" or not parsed.fragment.startswith("/oauth/authorize?") or
                        not re.fullmatch(r"[a-f0-9]{64}", str(result.get("flow_id", "")))):
                    raise ApolloError("The Apollo sign-in response was invalid.", "APOLLO_SERVICE", 502)
                self._write({"broker_url":self.client.origin,"flow":{"id":result["flow_id"], "secret":secret, "url":auth, "expires":self.clock()+600}})
                return {"authorization_url":auth, "connecting":True}
            if action == "poll":
                flow = data.get("flow", {})
                if not flow or flow.get("expires", 0) < self.clock():
                    raise ApolloError("Apollo sign-in expired. Connect again.", "APOLLO_CONNECT", 409)
                result = self.client.post("poll", {"flow_id":flow["id"], "secret":flow["secret"]})
                if result.get("state") == "pending":
                    return {"connecting":True}
                if result.get("state") != "connected" or result.get("connection_id") != flow["id"]:
                    self._write({})
                    raise ApolloError("Apollo sign-in was declined or interrupted. Check your account permissions and connect again.", "APOLLO_CONNECT", 409)
                name = result.get("name")
                if not isinstance(name, str) or not name.strip() or len(name) > 160 or any(ord(c)<32 for c in name):
                    raise ApolloError("Apollo did not confirm the account identity.", "APOLLO_SERVICE", 502)
                self._write({"broker_url":self.client.origin,"connection":flow["id"], "secret":flow["secret"], "name":name, "emails":{}})
                return {"connected":True, "name":name}
            if action == "disconnect":
                # Disconnect must remain possible even if the ad database or
                # contact archive is unavailable. Its confirmed cache clearing
                # remains separate from already archived device contacts.
                warning = ""
                try:
                    if data.get("connection"):
                        self.client.post("disconnect", self._credentials(data))
                    elif data.get("flow"):
                        self.client.post("cancel", {"flow_id":data["flow"]["id"], "secret":data["flow"]["secret"]})
                except ApolloError:
                    warning = "Disconnected on this device. Remove CV Studio in Apollo's connected integrations to revoke any remaining authorization."
                # Overwrite authoritatively rather than relying on best-effort
                # file deletion (which could expose an older fallback vault).
                self._write({})
                return {"connected":False, "warning":warning}
            credentials = self._credentials(data)
            if not isinstance(row, dict) or row.get("deleted") or not row.get("hr"):
                raise ApolloError("Open a saved ad with saved HR people first.")
            people = row["hr"]["people"]
            def identity(person):
                return self._identity(row, person)
            cache = data.setdefault("emails", {})
            if action == "emails":
                warning = self._archive_warning(row, cache)
                return {"emails":{p["url"]:self._public(cache[identity(p)]) for p in people if identity(p) in cache}, "contact_warning":warning}
            if action != "reveal":
                raise ApolloError("Choose a supported Apollo action.")
            person = next((p for p in people if p["url"].rstrip("/") == str(url).rstrip("/")), None)
            if not person:
                raise ApolloError("This person is no longer in the saved HR results. Refresh the ad.", "APOLLO_CONFLICT", 409)
            key = identity(person)
            previous = cache.get(key)
            if previous and previous["state"] not in {"pending", "unconfirmed"}:
                return {"result":self._public(previous), "contact_warning":self._archive_warning(row, cache)}
            if previous:
                # Read-only recovery: the same operation is never resubmitted.
                action, body = "lookup-status", {**credentials, "operation_id":previous["operation_id"]}
            else:
                if confirmed is not True:
                    raise ApolloError("Confirm Apollo credit use for this email unlock.", "APOLLO_CONFIRM", 403)
                if len(cache) >= 500:
                    raise ApolloError("This connection has reached its local limit of 500 email lookups. Keep a copy of any emails you need before disconnecting.", "APOLLO_LIMIT", 409)
                previous = {"state":"pending", "operation_id":secrets.token_hex(32)}
                cache[key] = previous
                self._write(data)  # A durable non-replay marker precedes transport.
                action, body = "reveal", {**credentials, "confirmed":True, "operation_id":previous["operation_id"],
                    "person":{"url":person["url"], "name":person["name"], "company":row["company"],
                              "company_url":row["hr"]["company_url"], "company_id":row["hr"]["company_id"]}}
            try:
                result = self.client.post(action, body)
                state = result.get("state")
                if state not in RESULT_MESSAGES:
                    raise ValueError()
                email = result.get("email", "")
                email_status = result.get("email_status", "unknown")
                if state == "found":
                    if (not isinstance(email, str) or len(email) > 254 or not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,63}", email) or
                            email_status not in {"verified", "unverified", "catch_all", "unavailable", "user_managed", "unknown"}):
                        raise ValueError()
                previous.update(state=state, email=email if state == "found" else "", email_status=email_status if state == "found" else "unknown")
            except ApolloError as exc:
                previous.update(state="auth_required" if exc.code=="APOLLO_AUTH" else "unconfirmed", email="", email_status="unknown")
            except Exception:
                previous.update(state="unconfirmed", email="", email_status="unknown")
            self._write(data)
            return {"result":self._public(previous), "contact_warning":self._archive_warning(row, cache)}
        finally:
            self.lock.release()

    def _archive_warning(self, row, cache):
        try:
            self._archive([row], cache)
            return ""
        except ApolloError as exc:
            # The result/non-replay marker is already durable. Do not hide it or
            # resend the paid lookup just because the contact archive is full.
            return str(exc)

    @staticmethod
    def _public(result):
        state = result["state"]
        return {"state":state, "email":result.get("email", ""), "email_status":result.get("email_status", "unknown"),
                "message":RESULT_MESSAGES[state]}


class ApolloContacts:
    """Device-local contact history, separate from revocable OAuth credentials.

    Inject the existing protected vault; never store business emails in SQLite
    or browser storage. Source metadata is a snapshot, not a claim of freshness.
    """
    MAX_CONTACTS = 5000
    MAX_BYTES = 8 * 1024 * 1024
    FIELDS = {"name":200, "position":500, "company":200, "country":100,
              "location":240, "email":254, "email_status":32, "profile_url":400,
              "company_id":24, "company_url":300, "job_url":2048, "unlocked_at":40}

    def __init__(self, *, load, save, clock=time.time):
        self.load, self.save, self.clock = load, save, clock

    def _read(self):
        try:
            data = self.load()
            if not data:
                return {}
            if data.get("version") != 1 or not isinstance(data.get("contacts"), dict):
                raise ValueError()
            contacts = data["contacts"]
            if len(contacts) > self.MAX_CONTACTS:
                raise ValueError()
            for key, row in contacts.items():
                if not re.fullmatch(r"[a-f0-9]{64}", key) or not isinstance(row, dict) or set(row) != set(self.FIELDS):
                    raise ValueError()
                if any(not isinstance(row[k], str) or len(row[k]) > n or any(ord(c)<32 for c in row[k]) for k,n in self.FIELDS.items()):
                    raise ValueError()
            return contacts
        except Exception:
            raise ApolloError("Saved contacts could not be read securely. Keep the existing files and check local storage.", "APOLLO_CONTACTS_STORAGE", 503) from None

    def remember(self, entries):
        found = [(row, person, result) for row,person,result in entries if result.get("state") == "found"]
        if not found:
            return
        contacts = self._read()
        changed = False
        for row, person, result in found:
            # Deduplicate across job ads/accounts on this device, retaining the
            # first observed context for the same person + employer + email.
            key = hashlib.sha256(json.dumps([row["hr"]["company_id"], person["url"].rstrip("/").lower(), result["email"].lower()]).encode()).hexdigest()
            if key in contacts:
                continue
            contacts[key] = {"name":person["name"], "position":person.get("title", ""), "company":row["company"],
                "country":row["hr"].get("country", ""), "location":person.get("location", ""),
                "email":result["email"], "email_status":result.get("email_status", "unknown"),
                "profile_url":person["url"], "company_id":row["hr"]["company_id"],
                "company_url":row["hr"]["company_url"], "job_url":row.get("url", ""),
                "unlocked_at":datetime.fromtimestamp(self.clock(), timezone.utc).isoformat(timespec="seconds")}
            changed = True
        if not changed:
            return
        value = {"version":1, "contacts":contacts}
        if len(contacts) > self.MAX_CONTACTS or len(json.dumps(value).encode()) > self.MAX_BYTES:
            raise ApolloError("The contacts dashboard reached its local limit. This email is still saved on its HR card. Export your contacts and keep a copy of this email.", "APOLLO_CONTACTS_LIMIT", 409)
        try:
            self.save(value)
        except Exception:
            raise ApolloError("The email is saved on its HR card, but the contacts dashboard could not save it. Check storage, then refresh contacts; no credit will be spent again.", "APOLLO_CONTACTS_STORAGE", 503) from None

    def list(self, query=""):
        if not isinstance(query, str) or len(query) > 240 or any(ord(c)<32 for c in query):
            raise ApolloError("Use a shorter contacts search.")
        # Match the browser's toLowerCase search so exports contain exactly the
        # same matching rows (casefold would also expand e.g. ß into ss).
        query = query.strip().lower()
        return sorted(({**row,"contact_key":key} for key,row in self._read().items()
            if not query or query in " ".join(row.values()).lower()), key=lambda row:row["unlocked_at"], reverse=True)

    def get(self, key):
        if not isinstance(key,str) or not re.fullmatch(r"[a-f0-9]{64}",key):
            raise ApolloError("Choose a saved unlocked contact.")
        row=self._read().get(key)
        if not row:
            raise ApolloError("This saved contact is no longer available. Refresh contacts.", "APOLLO_CONTACT_CONFLICT", 409)
        return dict(row)


def export_contacts(rows):
    """A genuine XLSX file. Every source cell is literal text, never a formula."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    fields = {"name":"Name", "position":"Position", "company":"Company", "country":"Country (scan filter)",
              "location":"Observed location", "email":"Email", "email_status":"Apollo email status",
              "profile_url":"LinkedIn profile", "company_url":"LinkedIn company", "job_url":"Source job ad", "unlocked_at":"Saved at (UTC)"}
    book = Workbook()
    sheet = book.active
    sheet.title = "Unlocked contacts"
    sheet.append(list(fields.values()))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16365D")
    for number, row in enumerate(rows, 2):
        for column, key in enumerate(fields, 1):
            cell = sheet.cell(number, column, row.get(key, ""))
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1,len(fields)+1):
        sheet.column_dimensions[get_column_letter(column)].width = 30 if column != 2 else 45
    output = BytesIO()
    book.save(output)
    return output.getvalue()
