import copy
from io import BytesIO
from types import SimpleNamespace
import pytest
from unittest import mock
from openpyxl import load_workbook
from cvstudio_people import (SourcerCandidates,PeopleError,hr_people,filter_candidates,
    filter_people,workbook,CANDIDATE_FIELDS,employment_label)
from cvstudio_job_leads import JobLeadsStore,JobLeadsError
from test_job_leads_hr import AD,SEARCH


@pytest.fixture
def archive():
    data={}
    def save(value):
        data.clear();data.update(copy.deepcopy(value))
    store=SourcerCandidates(load=lambda:copy.deepcopy(data),save=save)
    return store,data


def candidate(**patch):
    return dict(name='Example Person',url='https://www.linkedin.com/in/example',title='SAP Consultant',
        location='Malaysia',score=87,evidence=['SAP Consultant'],unknowns=['Years not shown'],**patch)


def test_candidate_archive_persists_search_and_notes(archive):
    store,data=archive;run=SimpleNamespace(id='a'*32,brief='SAP consultant in Malaysia')
    store.remember(run,[candidate()], 'screen');row=store.list()[0]
    saved=store.note(row['key'],row['revision'],'=literal note')
    store.remember(run,[candidate()], 'detail');row=store.list()[0]
    assert row['notes']=='=literal note' and row['read']=='Profile detail'
    assert row['revision']==saved['revision']+1
    store.remember(run,[candidate()], 'screen')
    assert store.list()[0]['read']=='Profile detail'
    run.id='b'*32;store.remember(run,[candidate()], 'screen')
    assert len(store.list())==2
    assert len(filter_candidates(store.list(),'SAP',run.id))==1
    assert len(SourcerCandidates(load=lambda:copy.deepcopy(data),save=lambda _:None).list())==2
    with pytest.raises(PeopleError,match='changed'):store.note(row['key'],1,'stale')


def test_archive_failures_preserve_existing_data(archive):
    store,data=archive;run=SimpleNamespace(id='a'*32,brief='SAP')
    store.remember(run,[candidate()], 'screen');old=copy.deepcopy(data)
    store.MAX_ROWS=1;run.id='b'*32
    with pytest.raises(PeopleError,match='full'):store.remember(run,[candidate()],'detail')
    assert data==old
    run.brief='bad\x00brief'
    with pytest.raises(PeopleError,match='unsupported'):store.remember(run,[candidate()],'screen')
    assert data==old
    broken=SourcerCandidates(load=lambda:{'version':7},save=lambda _:pytest.fail('must not overwrite corruption'))
    with pytest.raises(PeopleError,match='read securely'):broken.remember(run,[candidate()],'screen')


def test_exports_are_literal_and_filter_matches(archive):
    store,_=archive;run=SimpleNamespace(id='a'*32,brief='=SUM(1)')
    store.remember(run,[candidate()],'screen');row=store.list()[0]
    store.note(row['key'],row['revision'],'@literal')
    rows=filter_candidates(store.list(),'@literal')
    book=load_workbook(BytesIO(workbook(rows,CANDIDATE_FIELDS,'Candidates')))
    sheet=book.active
    assert sheet.freeze_panes=='B2' and sheet.max_row==2
    assert all(cell.data_type!='f' for row in sheet for cell in row)
    assert sheet['I2'].value=='@literal'
    sanitized=load_workbook(BytesIO(workbook([{'notes':'text\x01with\nline'}],{'notes':'Notes'},'Candidates'))).active
    assert sanitized['A2'].value=='textwith\nline'


def test_hr_projection_excludes_trash_and_deduplicates(tmp_path):
    store=JobLeadsStore(tmp_path/'leads.sqlite3')
    row=store.save(AD)['record'];row=store.save_hr(row['id'],row['revision'],SEARCH)['record']
    duplicate={**copy.deepcopy(row),'id':'b'*32,'url':'https://www.linkedin.com/jobs/view/12345'}
    trash={**copy.deepcopy(row),'id':'c'*32,'deleted':True}
    people=hr_people([row,duplicate,trash])
    assert len(people)==1 and people[0]['ads']==2
    assert people[0]['id']==row['id'] and people[0]['revision']==row['revision']
    assert len(filter_people(people,'Example',people[0]['company_id']))==1
    assert not filter_people(people,'missing')
    result=hr_people([row,duplicate],lambda r,p:{'state':'found','email':'person@example.test'} if r['id']==duplicate['id'] else None)[0]
    assert result['email_id']==duplicate['id'] and result['email_result']['state']=='found'


def test_date_update_is_atomic_and_preserves_scan_metadata(tmp_path):
    store=JobLeadsStore(tmp_path/'leads.sqlite3');row=store.save(AD)['record']
    row=store.save_hr(row['id'],row['revision'],SEARCH)['record'];hr=copy.deepcopy(row['hr']);url=hr['people'][0]['url']
    data=store.save_hr(row['id'],row['revision'],None,person_url=url,dates={'joined':'2026-10','left':'','current':True})['record']
    assert employment_label(data['hr']['people'][0])=='Oct 2026 to Present'
    assert data['hr']['people'][0]['employment_status']=='observed'
    assert data['hr']['scanned_at']==hr['scanned_at'] and data['hr']['complete']==hr['complete']
    with pytest.raises(JobLeadsError):store.save_hr(row['id'],row['revision'],None,person_url=url,dates={'joined':'2020','left':'2024','current':False})
    with pytest.raises(JobLeadsError):store.save_hr(data['id'],data['revision'],None,person_url=url,dates={'joined':'2026-13','left':'','current':True})
    assert store.hr_record(data['id'],data['revision'])['hr']==data['hr']


@pytest.mark.parametrize('dates,label',[
    ({'joined':'2026-10','left':'','current':True},'Oct 2026 to Present'),
    ({'joined':'2020-06','left':'2024-03','current':False},'Jun 2020 to Mar 2024'),
    ({'joined':'2020'},'2020'), ({'joined':''},'Not checked'),
    ({'joined':'','employment_status':'not_shown'},'Not confirmed')])
def test_date_labels(dates,label):
    assert employment_label(dates)==label


def test_guarded_dashboards_exports_and_new_vault(archive,tmp_path):
    import app
    from test_apollo_contacts import setup
    store,_=archive;store.remember(SimpleNamespace(id='a'*32,brief='SAP Malaysia'),[candidate()],'screen')
    service,broker,*_=setup();before=len(broker.calls)
    existing=service.rows()[0]
    full={**existing,'url':'https://www.linkedin.com/jobs/view/12345','hr':{**existing['hr'],'country':'Malaysia','scanned_at':'2026-10-29T00:00:00Z'}}
    service.rows=lambda:[full]
    client=app.app.test_client();client.get('/');headers={'X-CV-Studio-Request':'1'}
    from cvstudio_linkedin_sourcer import LinkedInSourcerService
    sourcer=LinkedInSourcerService(key_available=lambda:False,provider_call=lambda *_:pytest.fail('No provider call'),candidates=store)
    with mock.patch.object(app,'_LINKEDIN_SOURCER',sourcer),mock.patch.object(app,'_SOURCER_CANDIDATES',store),mock.patch.object(app,'_APOLLO',service):
        assert client.post('/linkedin-sourcer',json={'action':'candidates'},headers={'Origin':'https://other.test'}).status_code==403
        response=client.post('/linkedin-sourcer',json={'action':'candidates'},headers=headers)
        assert response.status_code==200 and len(response.json['candidates'])==1
        assert client.post('/linkedin-sourcer',json={'action':'export_candidates'},headers=headers).data[:2]==b'PK'
        response=client.post('/apollo',json={'action':'hr_people'},headers=headers)
        assert response.status_code==200 and len(response.json['people'])==1
        assert 'no-store' in response.headers['Cache-Control']
        assert client.post('/apollo',json={'action':'export_people'},headers=headers).data[:2]==b'PK'
    assert len(broker.calls)==before
    paths=(str(tmp_path/'candidates.dpapi'),str(tmp_path/'candidates.enc'))
    with mock.patch.object(app,'_cv_secure_paths',return_value=paths):
        app._cv_secure_save('sourcer_candidates',{'version':1,'rows':{}})
        assert app._cv_secure_load('sourcer_candidates')['version']==1
        from pathlib import Path
        existing=next(p for p in map(Path,paths) if p.exists());assert b'"rows"' not in existing.read_bytes()
        existing.write_bytes(b'corrupt archive')
        with pytest.raises(ValueError):app._cv_secure_load('sourcer_candidates')


def test_paid_success_survives_archive_failure_and_metadata_stays_observed():
    from cvstudio_linkedin_sourcer import LinkedInSourcerService
    archive=SimpleNamespace(remember=mock.Mock(side_effect=PeopleError('Dashboard unavailable')))
    def provider(*_):
        return {'content':[{'type':'tool_use','name':'return_profile_review','input':{'candidates':[{'id':'p0','score':80,'evidence':['SAP Consultant'],'unknowns':[]}]}}]}, {'usage':{'input_tokens':10,'output_tokens':20}}
    service=LinkedInSourcerService(key_available=lambda:True,provider_call=provider,candidates=archive)
    run=service.start({'brief':'SAP Malaysia','mode':'standard','max_pages':1,'max_calls':2,'confirm_paid':True})['run']
    result=service.handle({'action':'review','run_id':run['id'],'stage':'screen','page':1,'profiles':[{'id':'p0','name':'Example','url':'https://www.linkedin.com/in/example','text':'SAP Consultant\nMalaysia','title':'SAP Consultant','location':'Invented'}]})
    assert result['dashboard_warning']=='Dashboard unavailable'
    assert result['usage']['input_tokens']==10 and result['candidates'][0]['title']=='SAP Consultant'
    assert 'location' not in result['candidates'][0]
    assert result['run']['calls']==1


@pytest.mark.parametrize('status,joined',[
    ('not_checked',''),('not_shown',''),('observed','2022-07')])
def test_employment_evidence_status_survives_reopen_and_backup(tmp_path,status,joined):
    store=JobLeadsStore(tmp_path/'leads.sqlite3');row=store.save(AD)['record']
    search=copy.deepcopy(SEARCH);search['people'][0].update(joined=joined,employment_status=status)
    saved=store.save_hr(row['id'],row['revision'],search)['record']
    reopened=JobLeadsStore(tmp_path/'leads.sqlite3').list()[0]
    assert reopened['hr']==saved['hr']
    restored=JobLeadsStore(tmp_path/'restored.sqlite3');restored.restore(store.backup())
    assert restored.list()[0]['hr']==saved['hr']


@pytest.mark.parametrize('status,joined',[
    ('observed',''),('not_checked','2022-07'),('not_shown','2022-07'),
    ('bogus',''),(None,''),([],''),({},''),(True,'')])
def test_bad_employment_status_cannot_overwrite_saved_people(tmp_path,status,joined):
    store=JobLeadsStore(tmp_path/'leads.sqlite3');row=store.save(AD)['record']
    row=store.save_hr(row['id'],row['revision'],SEARCH)['record']
    search=copy.deepcopy(SEARCH);search['people'][0].update(joined=joined,employment_status=status)
    with pytest.raises(JobLeadsError):store.save_hr(row['id'],row['revision'],search)
    assert store.list()[0]==row


def test_refreshed_older_ad_dates_survive_duplicate_projection_and_excel(tmp_path):
    store=JobLeadsStore(tmp_path/'leads.sqlite3')
    with mock.patch('cvstudio_job_leads._now',return_value='2026-10-01T00:00:00Z'):
        old=store.save(AD)['record'];old=store.save_hr(old['id'],old['revision'],SEARCH)['record']
    with mock.patch('cvstudio_job_leads._now',return_value='2026-10-02T00:00:00Z'):
        newer=store.save({**AD,'url':'https://www.linkedin.com/jobs/view/123456'})['record']
        newer=store.save_hr(newer['id'],newer['revision'],SEARCH)['record']
    person=old['hr']['people'][0]
    old=store.save_hr(old['id'],old['revision'],None,person_url=person['url'],dates={'joined':'2022-07','left':'','current':True})['record']
    assert old['hr']['scanned_at']=='2026-10-01T00:00:00Z'
    projected=hr_people(JobLeadsStore(tmp_path/'leads.sqlite3').list())[0]
    assert projected['ads']==2 and projected['id']==newer['id']
    assert projected['employment']=='Jul 2022 to Present'
    assert projected['employment_id']==old['id'] and projected['employment_revision']==old['revision']
    sheet=load_workbook(BytesIO(workbook([projected],{'employment':'Employment'},'People'))).active
    assert sheet['A2'].value=='Jul 2022 to Present'
    newer=store.save_hr(newer['id'],newer['revision'],None,person_url=person['url'],dates={'joined':'2024-03','left':'2026-08','current':False})['record']
    latest=hr_people(store.list())[0]
    assert latest['employment']=='Mar 2024 to Aug 2026' and latest['employment_id']==newer['id']
    restored=JobLeadsStore(tmp_path/'restored.sqlite3');restored.restore(store.backup())
    assert hr_people(restored.list())[0]['employment']==latest['employment']


@pytest.mark.parametrize('stamp',[None,'','not a date','2026-10-29T00:00:00',[],True])
def test_bad_employment_observation_time_is_atomic(tmp_path,stamp):
    store=JobLeadsStore(tmp_path/'leads.sqlite3');row=store.save(AD)['record']
    search=copy.deepcopy(SEARCH);search['people'][0].update(joined='2022-07',employment_checked_at=stamp)
    with pytest.raises(JobLeadsError):store.save_hr(row['id'],row['revision'],search)
    assert store.list()[0]==row
