"""The candidate's pay never reaches the generated Summary box.

A generated summary is written from the whole CV, and a CV can state the
candidate's current and expected salary. The one filter
(_cv_strip_pay_from_summary in cvstudio_cv_normalize.py) runs where each summary
is made -- /generate-ai for the CV Summary callers, /blind for a promoted source
summary -- so the preview shows exactly what the Word file will carry, and
/generate-docx runs it again on both of its paths as a last net. The rule itself
is pinned by the "summary_salary" cases in tests/fixtures/cv_guardrail_cases.json.
All data is synthetic.
"""

import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import zipfile
from lxml import etree

from owner_build_tools.build_protected import write_test_receipt

ROOT = Path(__file__).resolve().parents[1]
_MODULE_TEMPORARY = tempfile.TemporaryDirectory(prefix="cvstudio-summary-salary-")
with mock.patch.dict(os.environ, {
    "LOCALAPPDATA": _MODULE_TEMPORARY.name, "APPDATA": _MODULE_TEMPORARY.name,
    "CVSTUDIO_DB_PATH": str(Path(_MODULE_TEMPORARY.name) / "state" / "cv_studio.sqlite3"),
    "CVSTUDIO_STATE_DIR": str(Path(_MODULE_TEMPORARY.name) / "state"),
    "CVSTUDIO_JOB_STATE_PATH": str(Path(_MODULE_TEMPORARY.name) / "jobs.json"),
    "SALARY_COMPARISON_DATA_DIR": str(Path(_MODULE_TEMPORARY.name) / "salary"),
}), mock.patch("pathlib.Path.home", return_value=Path(_MODULE_TEMPORARY.name)), \
        mock.patch("os.path.expanduser", side_effect=lambda path: (
            str(Path(_MODULE_TEMPORARY.name) / path[2:]) if path.startswith("~/") else path)):
    write_test_receipt(ROOT, environment={
        "HOME": _MODULE_TEMPORARY.name, "LOCALAPPDATA": _MODULE_TEMPORARY.name,
    })
    import app

HEADERS = {"Origin": "http://127.0.0.1:5000", "X-CV-Studio-Request": "1"}

SUMMARY = [
    "**Senior Data Engineer** with 11 years in banking and gaming.",
    "Built streaming pipelines on Apache Flink. Expected salary RM16,000 per month.",
    "Current salary of MYR 12,000; salary expectations are negotiable.",
    "Expertise in compensation and benefits reporting for 500 staff.",
]
PAY_TEXT = ("RM16,000", "MYR 12,000", "Expected salary", "salary expectations")


def _cv(summary_bullets):
    return {
        "candidate": {"name": "Test Candidate", "current_company": "Northwind Data Sdn Bhd",
                      "current_position": "Senior Data Engineer"},
        "summary_bullets": summary_bullets,
        "work_experiences": [
            {"date_range": "Jun 2025 to Present", "company": "Northwind Data Sdn Bhd",
             "roles": [{"title": "Senior Data Engineer", "bullets": ["Built batch pipelines."]}]},
        ],
        "education": [], "certifications": [], "skills": [],
    }


def _document_xml(docx_bytes):
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


class SummarySalaryDocxTests(unittest.TestCase):
    def setUp(self):
        self.client = app.app.test_client()

    def _format(self, summary_bullets):
        response = self.client.post("/generate-docx", json={"data": _cv(summary_bullets)}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        return response.data

    def test_body_and_table_word_lists_survive_extraction_and_blind_parse(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        expected = ["Led delivery teams.", "Designed data platforms."]
        for in_table in (False, True):
            with self.subTest(in_table=in_table):
                doc = Document()
                doc.add_paragraph("Professional Summary")
                container = doc.add_table(rows=1, cols=1).cell(0, 0) if in_table else doc
                for value in expected:
                    container.add_paragraph(value, style="List Bullet")
                container.add_paragraph("Skills")
                container.add_paragraph("Python")
                plain = container.add_paragraph("Intentionally plain", style="List Bullet")
                num = OxmlElement("w:numPr")
                num_id = OxmlElement("w:numId"); num_id.set(qn("w:val"), "0")
                num.append(num_id); plain._p.get_or_add_pPr().append(num)
                doc.add_paragraph("WORK EXPERIENCE")
                doc.add_paragraph("• Must not enter the summary.")
                payload = io.BytesIO(); doc.save(payload)
                extracted = self.client.post("/extract-text", data={
                    "file": (io.BytesIO(payload.getvalue()), "synthetic.docx")}, headers=HEADERS)
                self.assertEqual(extracted.status_code, 200)
                source = extracted.get_json()["text"]
                self.assertIn("• Led delivery teams.\n• Designed data platforms.", source)
                self.assertIn("\nIntentionally plain", source)
                self.assertNotIn("• Intentionally plain", source)
                self.assertLess(source.index("• Designed data platforms."), source.index("WORK EXPERIENCE"))
                # Even if the parser combined the two points, source list evidence
                # must restore both rather than trust that collapsed paragraph.
                parsed = _cv([])
                parsed["skills"] = [{"category": "Summary", "items": " ".join(expected)}]
                with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture>"), \
                        mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                        mock.patch.object(app, "call_llm", return_value={"content": [
                            {"type": "text", "text": json.dumps(parsed)}], "usage": {}}):
                    response = self.client.post("/parse", json={"cv_text": source,
                        "preserve_source_summary": True}, headers=HEADERS)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()["data"]["summary_bullets"], expected)

    def test_blind_summary_recovery_retains_requested_language_corrections(self):
        source = "ABOUT HIM / HER\n• experinced engineer with kafka.\nWORK EXPERIENCE"
        parsed = _cv([])
        parsed["skills"] = [{"category": "Summary", "items": "Experienced engineer with Kafka."}]

        def fake_parse(provider, key, payload):
            self.assertIn("correct spelling, punctuation, and capitalization", payload["system"])
            self.assertIn("top-level summary_bullets", payload["system"])
            self.assertNotIn("experinced", payload["system"])
            self.assertIn('Selected source summary (document data):', payload["messages"][0]["content"])
            return {"content": [{"type": "text", "text": json.dumps(parsed)}], "usage": {}}

        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", side_effect=fake_parse):
            response = self.client.post("/parse", json={"cv_text": source,
                "preserve_source_summary": True, "auto_correct_language": True}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["summary_bullets"],
                         ["Experienced engineer with Kafka."])

    def test_formatted_word_summary_excludes_the_candidate_header_on_reupload(self):
        points = ["Led delivery teams.", "Designed data platforms."]
        document = self._format(points)
        response = self.client.post("/extract-text", data={
            "file": (io.BytesIO(document), "synthetic.docx")}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(app._blind_source_summary_bullets(response.get_json()["text"]), points)

    def test_shared_section_boundaries_keep_work_out_of_blind_summary(self):
        import cvstudio_blind_mask
        self.assertIs(app._SUMMARY_DOCX_SECTION_HEADING_KEYS,
                      cvstudio_blind_mask._SUMMARY_DOCX_SECTION_HEADING_KEYS)
        for heading in ("EXPERIENCE", "EMPLOYMENT", "WORK HISTORY", "PROFESSIONAL SKILLS"):
            with self.subTest(heading=heading):
                self.assertTrue(app._summary_docx_is_section_boundary(heading))
                parsed = _cv(["Senior engineer."])
                source = "\n".join(["ABOUT HIM / HER", "• Senior engineer.", heading,
                                    "• Built batch pipelines.", "Education"])
                with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture>"), \
                        mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                        mock.patch.object(app, "call_llm", return_value={"content": [
                            {"type": "text", "text": json.dumps(parsed)}], "usage": {}}):
                    response = self.client.post("/parse", json={"cv_text": source,
                        "preserve_source_summary": True}, headers=HEADERS)
                self.assertEqual(response.status_code, 200)
                result = response.get_json()["data"]
                self.assertEqual(result["summary_bullets"], ["Senior engineer."])
                self.assertEqual(result["work_experiences"][0]["roles"][0]["bullets"],
                                 ["Built batch pipelines."])

    def test_native_word_numbered_section_headings_remain_boundaries(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Professional Summary", style="List Number")
        doc.add_paragraph("Led delivery teams.", style="List Bullet")
        doc.add_paragraph("Designed data platforms.", style="List Bullet")
        doc.add_paragraph("Work Experience", style="List Number")
        doc.add_paragraph("Built pipelines.", style="List Bullet")
        payload = io.BytesIO(); doc.save(payload)
        response = self.client.post("/extract-text", data={
            "file": (io.BytesIO(payload.getvalue()), "synthetic.docx")}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        source = response.get_json()["text"]
        self.assertEqual(app._blind_source_summary_bullets(source),
                         ["Led delivery teams.", "Designed data platforms."])

    def test_requested_corrections_survive_a_source_pay_point_omitted_by_parser(self):
        source = "ABOUT HIM / HER\n• experinced engineer with kafka.\n• Current salary RM 12k.\nWORK EXPERIENCE"
        parsed = _cv([])
        parsed["skills"] = [{"category": "Summary", "items": "Experienced engineer with Kafka."}]
        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", return_value={"content": [
                    {"type": "text", "text": json.dumps(parsed)}], "usage": {}}):
            response = self.client.post("/parse", json={"cv_text": source,
                "preserve_source_summary": True, "auto_correct_language": True}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["data"]["summary_bullets"],
                         ["Experienced engineer with Kafka."])
        self.assertEqual(response.get_json()["summary_pay_removed"], 1)

    def test_format_cv_never_writes_pay_into_the_summary_box(self):
        xml = _document_xml(self._format(SUMMARY))
        for text in PAY_TEXT:
            self.assertNotIn(text, xml)
        # Everything that is not pay is kept, including the rest of a bullet that
        # also stated pay, and HR work that merely mentions compensation.
        self.assertIn("Senior Data Engineer", xml)
        self.assertIn("Built streaming pipelines on Apache Flink.", xml)
        self.assertIn("Expertise in compensation and benefits reporting for 500 staff.", xml)

    def test_explicit_earnings_leave_but_employee_pay_achievement_stays(self):
        work = "Managed payroll with salary of RM 12,000 for each employee."
        xml = _document_xml(self._format([work, "Earnings of RM 18,000 monthly.",
                                        "The candidate receives RM 18,000 per month."]))
        self.assertIn(work, xml)
        self.assertNotIn("18,000", xml)

    def test_work_opening_pay_is_removed_in_provider_summary_and_word(self):
        work = "Analysed company earnings of RM 24,000 monthly."
        pay = ["Led HR and the candidate receives RM 18,000 per month.",
               "Analysed performance and her earnings are RM 18,000 monthly."]
        out = self._generate_ai("\n".join("- " + text for text in [work] + pay),
                                strip_candidate_pay=True)
        self.assertEqual(out["summary_pay_removed"], 2)
        self.assertIn(work, out["content"][0]["text"])
        self.assertNotIn("18,000", out["content"][0]["text"])
        response = self.client.post("/generate-docx", json={"data": _cv([work] + pay)}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-CV-Summary-Pay-Removed"], "2")
        self.assertIn(work, _document_xml(response.data))
        self.assertNotIn("18,000", _document_xml(response.data))

    def test_the_uploaded_docx_summary_path_never_writes_pay(self):
        source = self._format(["Placeholder summary."])
        response = self.client.post(
            "/generate-docx",
            data={"source_docx": (io.BytesIO(source), "CV.docx"), "summary_bullets": json.dumps(SUMMARY)},
            content_type="multipart/form-data", headers={"Origin": "http://127.0.0.1:5000"},
        )
        self.assertEqual(response.status_code, 200)
        xml = _document_xml(response.data)
        for text in PAY_TEXT:
            self.assertNotIn(text, xml)
        self.assertIn("Built streaming pipelines on Apache Flink.", xml)
        self.assertIn("Expertise in compensation and benefits reporting for 500 staff.", xml)

    def test_the_uploaded_docx_path_reports_what_it_removed(self):
        source = self._format(["Placeholder summary."])
        response = self.client.post(
            "/generate-docx",
            data={"source_docx": (io.BytesIO(source), "CV.docx"),
                  "summary_bullets": json.dumps(["Built pipelines.", "Current salary RM 9k; expected RM 11k."])},
            content_type="multipart/form-data", headers={"Origin": "http://127.0.0.1:5000"},
        )
        self.assertEqual(response.status_code, 200)
        # One sentence, counted once however many clauses it had.
        self.assertEqual(response.headers.get("X-CV-Summary-Pay-Removed"), "1")
        clean = self.client.post(
            "/generate-docx",
            data={"source_docx": (io.BytesIO(source), "CV.docx"), "summary_bullets": json.dumps(["Built pipelines."])},
            content_type="multipart/form-data", headers={"Origin": "http://127.0.0.1:5000"},
        )
        self.assertEqual(clean.status_code, 200)
        self.assertNotIn("X-CV-Summary-Pay-Removed", clean.headers)

    def test_the_upload_helper_filters_once_and_keeps_its_old_shape(self):
        source = self._format(["Placeholder summary."])
        calls = []
        real = app._cv_strip_pay_from_summary_counted

        def counted(bullets):
            calls.append(1)
            return real(bullets)

        with mock.patch.object(app, "_cv_strip_pay_from_summary_counted", side_effect=counted):
            document, removed = app._insert_summary_into_docx_bytes_counted(
                source, ["Built pipelines.", "Expected salary RM 9k."])
        self.assertEqual((len(calls), removed), (1, 1))
        self.assertNotIn("RM 9k", _document_xml(document))
        # The original helper still returns the document alone.
        self.assertIsInstance(app._insert_summary_into_docx_bytes(source, ["Built pipelines."]), bytes)

    def test_the_filter_reads_only_what_the_word_file_can_use(self):
        seen = []
        real = app._cv_strip_pay_from_summary_counted

        def counted(bullets):
            seen.append(bullets)
            return real(bullets)

        huge = ["Built pipelines. " * 5000] + ["Led a team."] * 500
        with mock.patch.object(app, "_cv_strip_pay_from_summary_counted", side_effect=counted):
            response = self.client.post("/generate-docx", json={"data": _cv(huge)}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(seen[-1]), app._CV_SUMMARY_FILTER_MAX_ITEMS)
        self.assertLessEqual(max(len(item) for item in seen[-1]), app._CV_SUMMARY_FILTER_MAX_CHARS)
        # The uploaded-DOCX path reads no more than that either.
        source = self._format(["Placeholder summary."])
        with mock.patch.object(app, "_cv_strip_pay_from_summary_counted", side_effect=counted):
            app._insert_summary_into_docx_bytes_counted(source, huge)
        self.assertEqual(len(seen[-1]), app._CV_SUMMARY_FILTER_MAX_ITEMS)
        self.assertLessEqual(max(len(item) for item in seen[-1]), app._CV_SUMMARY_FILTER_MAX_CHARS)
        # A normal summary is passed through untouched.
        normal = ["Built pipelines.", "Led a team."]
        self.assertEqual(app._bounded_summary_bullets(normal), normal)
        self.assertEqual(app._bounded_summary_bullets([0, None, "x"]), [0, None, "x"])

    def test_a_summary_without_pay_is_written_unchanged(self):
        clean = ["**Senior Data Engineer** with 11 years in banking.", "Built streaming pipelines."]
        self.assertEqual(app._cv_strip_pay_from_summary(list(clean)), clean)
        xml = _document_xml(self._format(clean))
        self.assertIn("Built streaming pipelines.", xml)

    def test_a_summary_that_was_only_pay_says_so(self):
        source = self._format(["Placeholder summary."])
        response = self.client.post(
            "/generate-docx",
            data={"source_docx": (io.BytesIO(source), "CV.docx"),
                  "summary_bullets": json.dumps(["Expected salary RM16,000.", "Salary: RM 12k"])},
            content_type="multipart/form-data", headers={"Origin": "http://127.0.0.1:5000"},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("only stated the candidate's pay", response.get_json()["error"])
        # An empty request keeps its original message.
        empty = self.client.post(
            "/generate-docx",
            data={"source_docx": (io.BytesIO(source), "CV.docx"), "summary_bullets": "[]"},
            content_type="multipart/form-data", headers={"Origin": "http://127.0.0.1:5000"},
        )
        self.assertEqual(empty.get_json()["error"], "No CV Summary bullets provided")

    def _generate_ai(self, text, **extra):
        body = {"prompt": "fixture prompt", "provider": "anthropic"}
        body.update(extra)
        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", return_value={
                    "content": [{"type": "text", "text": text}], "usage": {}}):
            response = self.client.post("/generate-ai", json=body, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        return response.get_json()

    def test_generate_ai_filters_the_summary_when_asked(self):
        raw = "- **Senior engineer** with 11 years.\n- Expected salary RM16,000.\n- Built pipelines."
        out = self._generate_ai(raw, strip_candidate_pay=True)
        self.assertEqual(out["content"][0]["text"], "- **Senior engineer** with 11 years.\n- Built pipelines.")
        self.assertEqual(out["summary_pay_removed"], 1)
        # A clean summary is returned exactly as the provider wrote it.
        clean = self._generate_ai("- Built pipelines.", strip_candidate_pay=True)
        self.assertEqual(clean["content"], [{"type": "text", "text": "- Built pipelines."}])
        self.assertEqual(clean["summary_pay_removed"], 0)

    def test_generate_ai_is_unchanged_for_every_other_caller(self):
        # Blind JD, Company Profile and the rest do not ask, so nothing changes.
        raw = "Offered salary RM16,000 for the role."
        out = self._generate_ai(raw)
        self.assertEqual(out["content"][0]["text"], raw)
        self.assertNotIn("summary_pay_removed", out)
        out = self._generate_ai(raw, strip_candidate_pay="yes")
        self.assertEqual(out["content"][0]["text"], raw)

    def test_the_anonymized_summary_is_filtered_too(self):
        out = self._generate_ai(
            "- Built pipelines.\n- Expected salary RM16,000.",
            strip_candidate_pay=True, feature="summary_anonymized",
            source_cv_text="Test Candidate\nSalary Expectations: 16000",
        )
        self.assertNotIn("RM16,000", out["content"][0]["text"])
        self.assertIn("Built pipelines.", out["content"][0]["text"])

    def test_blind_parse_keeps_source_about_points_through_masking_and_word(self):
        source = "\n".join([
            "Test Candidate", "ABOUT HIM / HER",
            "• Engineer at Northwind Data Sdn Bhd with six years", "of experience.",
            "• Delivered migrations.", "• Current salary RM 12k.",
            "Note: Open to training.", "Private & Confidential | © Hyppies.com",
            "WORK EXPERIENCES", "Summary:", "A later generic paragraph.",
        ])
        parsed = _cv([])
        parsed["skills"] = [{"category": "Summary", "items": "A later generic paragraph."}]

        def fake_parse(provider, key, payload):
            return {"content": [{"type": "text", "text": json.dumps(parsed)}], "usage": {}}

        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", side_effect=fake_parse):
            for flag in (None, False, "true", True):
                body = {"cv_text": source, "provider": "anthropic"}
                if flag is not None:
                    body["preserve_source_summary"] = flag
                response = self.client.post("/parse", json=body, headers=HEADERS)
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
                if flag is not True:
                    self.assertEqual(response.get_json()["data"]["summary_bullets"], [])
            recovered = response.get_json()
        expected = ["Engineer at Northwind Data Sdn Bhd with six years of experience.",
                    "Delivered migrations.", "Note: Open to training."]
        self.assertEqual(recovered["data"]["summary_bullets"], expected)
        self.assertEqual(recovered["summary_pay_removed"], 1)
        seen = {}

        def fake_blind(provider, key, payload):
            sent = json.loads(payload["messages"][0]["content"].split("\n\n", 1)[1])
            seen["summary"] = sent["summary_bullets"][:]
            self.assertEqual(sent["skills"], [])
            # Echoing source identity is deliberately hostile: the existing
            # deterministic scrub still has to mask the recovered source too.
            return {"content": [{"type": "text", "text": json.dumps(sent)}], "usage": {}}

        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", side_effect=fake_blind):
            response = self.client.post("/blind", json={"cv_data": recovered["data"],
                                                       "provider": "anthropic"}, headers=HEADERS)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(seen["summary"], expected)
        blinded = response.get_json()["data"]
        self.assertEqual(len(blinded["summary_bullets"]), 3)
        self.assertNotIn("Northwind", json.dumps(blinded))
        self.assertNotIn("Test Candidate", json.dumps(blinded))
        self.assertNotIn("later generic", json.dumps(blinded))
        word = self.client.post("/generate-docx", json={"data": blinded}, headers=HEADERS)
        self.assertEqual(word.status_code, 200)
        xml = etree.fromstring(_document_xml(word.data).encode("utf-8"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        boxes = xml.xpath("//w:txbxContent", namespaces=ns)
        matching = [box for box in boxes if "Delivered migrations." in "".join(box.itertext())]
        self.assertEqual(len(matching), 2)  # DrawingML plus the VML compatibility copy.
        for box in matching:
            paragraphs = box.xpath("./w:p[w:pPr/w:numPr]", namespaces=ns)
            self.assertEqual(len(paragraphs), 3)
            self.assertEqual(["".join(p.xpath(".//w:t/text()", namespaces=ns))
                              for p in paragraphs], blinded["summary_bullets"])

    def test_the_blind_summary_is_filtered_before_the_provider_sees_it(self):
        seen = {}

        def fake_call(provider, key, payload):
            sent = json.loads(payload["messages"][0]["content"].split("\n\n", 1)[1])
            seen["summary"] = sent.get("summary_bullets")
            return {"content": [{"type": "text", "text": json.dumps(sent)}], "usage": {}}

        cv = _cv([])
        cv["skills"] = [{"category": "Summary",
                         "items": "Data engineer with 11 years in banking.\nSeeking a role with salary above 10000."}]
        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", side_effect=fake_call):
            response = self.client.post("/blind", json={"cv_data": cv, "provider": "anthropic"}, headers=HEADERS)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(seen["summary"], ["Data engineer with 11 years in banking."])
        # The preview the browser gets is the summary the Word file will carry,
        # and the removal is reported.
        self.assertEqual(response.get_json()["data"]["summary_bullets"], ["Data engineer with 11 years in banking."])
        self.assertEqual(response.get_json()["summary_pay_removed"], 1)

    def test_pay_only_recovered_about_cannot_fall_back_to_a_later_summary(self):
        source = "ABOUT HIM / HER\n• Current salary RM 12k.\nWORK EXPERIENCES\nSummary:\nLater paragraph."
        parsed = _cv([])
        parsed["skills"] = [{"category": "Summary", "items": "Later paragraph."},
                            {"category": "Technology", "items": "Python"}]
        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", return_value={"content": [
                    {"type": "text", "text": json.dumps(parsed)}], "usage": {}}):
            response = self.client.post("/parse", json={"cv_text": source,
                "preserve_source_summary": True}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["summary_pay_removed"], 1)
        self.assertEqual(result["data"]["summary_bullets"], [])
        self.assertEqual(app._blind_prepare_summary_bullets(result["data"])["summary_bullets"], [])
        self.assertEqual(result["data"]["skills"], [{"category": "Technology", "items": "Python"}])

    def test_a_format_summary_that_was_only_pay_still_writes_the_cv(self):
        # The summary is one part of the CV: the Word file is still written, with
        # an empty box, and the removal is reported so the page can say so.
        response = self.client.post(
            "/generate-docx", json={"data": _cv(["**Salary:** RM 17,000 per month.", "Expected salary RM 9k."])},
            headers=HEADERS,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers.get("X-CV-Summary-Pay-Removed"), "2")
        xml = _document_xml(response.data)
        self.assertNotIn("RM 17,000", xml)
        self.assertNotIn("RM 9k", xml)
        # Nothing removed, nothing reported.
        clean = self.client.post("/generate-docx", json={"data": _cv(["Built pipelines."])}, headers=HEADERS)
        self.assertEqual(clean.status_code, 200)
        self.assertNotIn("X-CV-Summary-Pay-Removed", clean.headers)
        self.assertEqual(self.client.post("/generate-docx", json={"data": _cv([])}, headers=HEADERS).status_code, 200)

    def test_the_parsed_summary_is_filtered_so_the_preview_matches_the_word_file(self):
        parsed = _cv(["Data engineer with 11 years in banking.", "Current salary of **RM 12,000** per month."])

        def fake_call(provider, key, payload):
            return {"content": [{"type": "text", "text": json.dumps(parsed)}], "usage": {}}

        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", side_effect=fake_call):
            response = self.client.post("/parse", json={"cv_text": "Test Candidate\nData engineer."}, headers=HEADERS)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["data"]["summary_bullets"], ["Data engineer with 11 years in banking."])
        self.assertEqual(response.get_json()["summary_pay_removed"], 1)

    def test_a_parse_with_nothing_removed_reports_nothing(self):
        parsed = _cv(["Data engineer with 11 years in banking."])

        def fake_call(provider, key, payload):
            return {"content": [{"type": "text", "text": json.dumps(parsed)}], "usage": {}}

        with mock.patch.object(app, "_resolve_request_api_key", return_value="<fixture-credential>"), \
                mock.patch.object(app, "_ai_spend_session_allowed", return_value=True), \
                mock.patch.object(app, "call_llm", side_effect=fake_call):
            response = self.client.post("/parse", json={"cv_text": "Test Candidate\nData engineer."}, headers=HEADERS)
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("summary_pay_removed", response.get_json())

    def test_an_intro_line_is_not_left_as_the_summary(self):
        for raw in ("Here is the summary:\n- Expected salary RM 9k",
                    "Here is the summary:\n1. Salary: RM 9,000\n2. Expected salary RM 10k"):
            with self.subTest(raw=raw):
                out = self._generate_ai(raw, strip_candidate_pay=True)
                self.assertEqual(out["content"], [{"type": "text", "text": ""}])
                self.assertGreater(out["summary_pay_removed"], 0)

    def test_the_summary_instructions_forbid_pay(self):
        source = (ROOT / "vendor" / "cvstudio" / "candidate-summary.js").read_text(encoding="utf-8")
        self.assertIn("Never mention the candidate\\'s salary or pay in any form", source)


if __name__ == "__main__":
    unittest.main()
