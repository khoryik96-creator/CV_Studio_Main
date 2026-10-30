"""Synthetic contact history: no credentials, network or live addresses."""
import copy
from io import BytesIO
from unittest import mock

import pytest
from openpyxl import load_workbook
from cvstudio_apollo import ApolloContacts, ApolloError, export_contacts
from test_apollo import fixture, connect, ROW, PERSON


def setup():
    service, broker, vault = fixture()
    archive = {}
    def save(value):
        archive.clear()
        archive.update(copy.deepcopy(value))
    contacts = ApolloContacts(load=lambda:copy.deepcopy(archive), save=save, clock=lambda:1000)
    service.contacts = contacts
    service.rows = lambda:[ROW]
    connect(service)
    return service, broker, vault, contacts, archive


def reveal(service, row=ROW):
    return service.perform('reveal', row=row, url=PERSON['url'], confirmed=True)


def test_found_contact_is_automatic_protected_and_survives_disconnect_and_restart():
    service, broker, vault, contacts, archive = setup()
    assert service.perform('contacts')['contacts'] == []
    result = reveal(service)
    assert result['result']['state']=='found' and not result['contact_warning']
    rows = service.perform('contacts')['contacts']
    assert len(rows)==1
    assert rows[0]['name']==PERSON['name'] and rows[0]['position']==PERSON['title']
    assert rows[0]['company']==ROW['company'] and rows[0]['email']=='person@example.test'
    assert rows[0]['country']=='' and rows[0]['location']=='Malaysia', 'do not invent a country from free-form location'
    assert not {'secret','connection','operation_id'} & rows[0].keys()
    before=len(broker.calls)
    assert service.perform('contacts', query='EXAMPLE')['contacts']==rows
    assert service.perform('contacts', query='unmatched')['contacts']==[]
    reveal(service)
    assert len(archive['contacts'])==1 and len(broker.calls)==before
    service.perform('disconnect')
    assert not vault and service.perform('contacts')['contacts']==rows
    restarted=ApolloContacts(load=lambda:copy.deepcopy(archive), save=mock.Mock())
    assert restarted.list()==rows
    connect(service)
    assert service.perform('emails',row=ROW)['emails']=={}, 'old account lookups do not leak to a new account'
    assert service.perform('contacts')['contacts']==rows, 'explicit device contact history remains available'


def test_same_person_across_ads_is_deduplicated_but_different_company_is_kept():
    service, broker, vault, contacts, archive=setup()
    reveal(service)
    second=copy.deepcopy(ROW);second['id']='2'*32
    reveal(service,second)
    assert len(contacts.list())==1
    second['hr']['company_id']='456';second['company']='Another Company'
    second['hr']['country']='Malaysia'
    reveal(service,second)
    assert len(contacts.list())==2 and contacts.list(query='Another')[0]['country']=='Malaysia'


def test_unicode_export_search_matches_browser_lowercase_without_casefold_expansion():
    service,broker,vault,contacts,archive=setup()
    row=copy.deepcopy(ROW);row['company']='Straße'
    reveal(service,row)
    assert len(contacts.list(query='ß'))==1
    assert contacts.list(query='STRASSE')==[]


@pytest.mark.parametrize('state',['not_found','mismatch','pending','unconfirmed','limited','auth_required'])
def test_unsuccessful_results_are_not_contacts(state):
    service, broker, vault, contacts, archive=setup()
    broker.result={'state':state}
    reveal(service)
    assert contacts.list()==[]


def test_archive_failure_keeps_paid_result_and_refresh_recovers_without_new_lookup():
    service, broker, vault, contacts, archive=setup()
    save=contacts.save
    contacts.save=mock.Mock(side_effect=OSError('fictional private path'))
    result=reveal(service)
    assert result['result']['state']=='found'
    assert 'dashboard could not save' in result['contact_warning'] and 'private' not in str(result)
    assert next(iter(vault['emails'].values()))['state']=='found'
    before=len(broker.calls)
    assert service.perform('contacts')['contact_warning']
    contacts.save=save
    assert len(service.perform('contacts')['contacts'])==1
    assert len(broker.calls)==before
    reveal(service)
    assert len(broker.calls)==before


def test_old_cached_results_are_imported_without_spend_and_bad_archive_is_not_overwritten():
    service, broker, vault, contacts, archive=setup()
    service.contacts=None
    reveal(service)
    service.contacts=contacts
    before=len(broker.calls)
    assert len(service.perform('contacts')['contacts'])==1
    assert len(broker.calls)==before
    archive['version']=999
    save=mock.Mock();contacts.save=save
    with pytest.raises(ApolloError,match='read securely'):
        service.perform('contacts')
    save.assert_not_called()


def test_disconnect_does_not_depend_on_ad_database_or_contact_archive_health():
    service,broker,vault,contacts,archive=setup()
    reveal(service)
    original=copy.deepcopy(archive)
    service.rows=mock.Mock(side_effect=OSError('fictional unavailable ad database'))
    contacts.load=mock.Mock(side_effect=ValueError('fictional archive problem'))
    service.perform('disconnect')
    assert not vault and broker.calls[-1][0]=='disconnect'
    service.rows.assert_not_called();contacts.load.assert_not_called()
    assert archive==original


def test_capacity_failure_is_atomic_and_keeps_hr_card_result():
    service, broker, vault, contacts, archive=setup()
    contacts.MAX_CONTACTS=0
    result=reveal(service)
    assert 'limit' in result['contact_warning'] and result['result']['email']=='person@example.test'
    assert not archive
    contacts.MAX_CONTACTS=5000
    assert len(service.perform('contacts')['contacts'])==1
    contacts.MAX_BYTES=1
    second=copy.deepcopy(ROW);second['hr']['company_id']='2'
    assert 'limit' in reveal(service,second)['contact_warning']
    assert len(archive['contacts'])==1
    service.rows=lambda:[ROW,second]
    response=service.perform('contacts')
    assert response['contacts'] and response['contact_warning'], 'capacity must not prevent export of saved history'


def test_failed_import_still_exports_saved_history_with_visible_incomplete_indicator():
    import app
    service,broker,vault,contacts,archive=setup()
    reveal(service)
    service.rows=mock.Mock(side_effect=OSError('fictional ad storage failure'))
    assert len(service.perform('contacts')['contacts'])==1
    assert service.perform('contacts')['contact_warning']
    client=app.app.test_client();client.get('/')
    before=len(broker.calls)
    with mock.patch.object(app,'_APOLLO',service):
        response=client.post('/apollo',json={'action':'export_contacts'},headers={'X-CV-Studio-Request':'1'})
        assert response.status_code==200 and response.headers['X-CVStudio-Contacts-Incomplete']=='1'
        assert load_workbook(BytesIO(response.data)).active['F2'].value=='person@example.test'
    assert len(broker.calls)==before


@pytest.mark.parametrize('query',[None,123,'x'*241,'bad\nquery'])
def test_invalid_search_rejected(query):
    service,broker,*_=setup()
    before=len(broker.calls)
    with pytest.raises(ApolloError):
        service.perform('contacts',query=query)
    assert len(broker.calls)==before


def test_real_xlsx_literal_cells_unicode_filters_and_no_credential_fields():
    row={'name':'=HYPERLINK("https://example.test","x")', 'company':'Éxample 中国',
         'position':'+cmd', 'country':'Malaysia', 'email':'person@example.test', 'email_status':'catch_all'}
    raw=export_contacts([row])
    assert raw.startswith(b'PK')
    book=load_workbook(BytesIO(raw))
    sheet=book.active
    assert sheet.title=='Unlocked contacts' and sheet.freeze_panes=='A2'
    assert sheet['A2'].value==row['name'] and sheet['A2'].data_type=='s'
    assert sheet['B2'].value=='+cmd' and sheet['B2'].data_type=='s'
    assert sheet['C2'].value=='Éxample 中国' and sheet['F2'].value==row['email']
    assert sheet.auto_filter.ref=='A1:K2'
    assert not any(cell.data_type=='f' for line in sheet for cell in line)


def test_guarded_dashboard_and_xlsx_routes_do_not_contact_a_provider_or_modify_ads():
    import app
    service,broker,vault,contacts,archive=setup()
    reveal(service)
    before=len(broker.calls)
    client=app.app.test_client();client.get('/')
    headers={'X-CV-Studio-Request':'1'}
    with mock.patch.object(app,'_APOLLO',service):
        assert client.post('/apollo',json={'action':'contacts'},headers={'Origin':'https://other.test'}).status_code==403
        response=client.post('/apollo',json={'action':'contacts'},headers=headers)
        assert len(response.json['contacts'])==1 and 'no-store' in response.headers['Cache-Control']
        response=client.post('/apollo',json={'action':'export_contacts','query':'Example'},headers=headers)
        assert response.status_code==200 and 'no-store' in response.headers['Cache-Control']
        assert 'CV-Studio-Contacts.xlsx' in response.headers['Content-Disposition']
        sheet=load_workbook(BytesIO(response.data)).active
        assert sheet.max_row==2 and sheet['F2'].value=='person@example.test'
        response=client.post('/apollo',json={'action':'export_contacts','query':'unmatched'},headers=headers)
        assert load_workbook(BytesIO(response.data)).active.max_row==1
    assert len(broker.calls)==before


def test_new_archive_uses_protected_files_and_detects_tampering(tmp_path):
    import app
    paths=(str(tmp_path/'contacts.dpapi'),str(tmp_path/'contacts.enc'))
    with mock.patch.object(app,'_cv_secure_paths',return_value=paths):
        value={'version':1,'contacts':{}}
        app._cv_secure_save('apollo_contacts',value)
        assert app._cv_secure_load('apollo_contacts')['version']==1
        path=next(path for path in map(__import__('pathlib').Path,paths) if path.exists())
        path.write_bytes(b'broken secure archive')
        with pytest.raises(ValueError,match='read securely'):
            app._cv_secure_load('apollo_contacts')


def test_contact_filters_intersect_and_export_only_the_matching_rows():
    service,broker,vault,contacts,archive=setup()
    for company,country,status,name in [('123','Malaysia','verified','Alpha'),('123','Singapore','unverified','Beta'),('456','Malaysia','verified','Gamma')]:
        row=copy.deepcopy(ROW);row['hr']['company_id']=company;row['hr']['country']=country
        person={**PERSON,'name':name,'url':'https://www.linkedin.com/in/'+name.lower()}
        contacts.remember([(row,person,{'state':'found','email':name.lower()+'@example.test','email_status':status})])
    before=len(broker.calls)
    rows=contacts.list(filters={'company':'123','country':'Malaysia','email_status':'verified'})
    assert [r['name'] for r in rows]==['Alpha']
    book=load_workbook(BytesIO(export_contacts(rows)))
    assert book.active.max_row==2 and book.active.cell(2,1).value=='Alpha'
    assert contacts.list('Gamma',filters={'company':'123'})==[]
    assert len(broker.calls)==before,'filters and export never unlock or call a provider'


@pytest.mark.parametrize('filters',[[],{'unknown':'x'},{'country':['Malaysia']},{'country':'x'*101},{'country':'bad\x00text'},{'company':'not-a-company-id'},{'email_status':'invented'}])
def test_contact_filters_reject_unsupported_values(filters):
    _,_,_,contacts,_=setup()
    with pytest.raises(ApolloError):contacts.list(filters=filters)


def test_local_dashboard_refresh_waits_for_an_overlapping_read_without_provider_work():
    import threading
    service,broker,vault,contacts,archive=setup();before=len(broker.calls)
    service.rows=lambda:[]
    started=threading.Event();finished=threading.Event();results=[]
    service.lock.acquire()
    def read():
        started.set()
        try:results.append(service.perform('hr_people'))
        finally:finished.set()
    worker=threading.Thread(target=read);worker.start();assert started.wait(2)
    try:
        assert not finished.wait(0.03),'local refresh waits instead of producing a normal UI busy error'
        with pytest.raises(ApolloError) as error:service.perform('reveal',row=ROW,url=PERSON['url'],confirmed=True)
        assert error.value.code=='APOLLO_BUSY','a paid operation still fails immediately rather than queuing'
    finally:service.lock.release()
    assert finished.wait(3);worker.join()
    assert results and len(broker.calls)==before
