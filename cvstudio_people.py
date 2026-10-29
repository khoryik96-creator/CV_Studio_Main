"""Bounded local people dashboards. No provider transport or credentials."""
from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import re
import threading
from urllib.parse import urlsplit, unquote


class PeopleError(Exception):
    pass


def workbook(rows, fields, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(list(fields.values()))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='16365D')
    for number, row in enumerate(rows, 2):
        for column, key in enumerate(fields, 1):
            value = row.get(key, '')
            if isinstance(value, list):
                value = '\n'.join(value)
            cell = sheet.cell(number, column, re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]', '', str(value)))
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    for column, key in enumerate(fields, 1):
        sheet.column_dimensions[get_column_letter(column)].width = 48 if key in {'evidence', 'unknowns', 'notes', 'brief'} else 28
    output = BytesIO()
    book.save(output)
    return output.getvalue()


CANDIDATE_FIELDS = {'name':'Name', 'title':'Headline / title', 'location':'Location',
    'url':'LinkedIn URL', 'score':'Evidence match / 100', 'read':'Profile read',
    'evidence':'Supporting evidence', 'unknowns':'Still unknown', 'notes':'Notes',
    'brief':'Search brief', 'saved_at':'Last assessed (UTC)'}
HR_FIELDS = {'name':'Name','title':'Position','company':'Company','location':'Observed location',
    'country':'Country (scan filter)','employment':'Employment','url':'LinkedIn URL',
    'matched':'Matched HR keywords','job_url':'Source ad','scanned_at':'Last scan (UTC)'}


def _filter(rows, query, selected, field, fields):
    if not isinstance(query,str) or len(query)>240 or not isinstance(selected,str) or len(selected)>240:
        raise PeopleError('Use a shorter dashboard search.')
    query=query.strip().lower()
    return [row for row in rows if (not selected or row.get(field)==selected) and
            (not query or query in ' '.join(' '.join(row.get(k,[])) if isinstance(row.get(k),list) else str(row.get(k,'')) for k in fields).lower())]


def filter_candidates(rows, query='', search=''):
    return _filter(rows,query,search,'run_id',CANDIDATE_FIELDS)


def filter_people(rows, query='', company=''):
    return _filter(rows,query,company,'company_id',HR_FIELDS)


def employment_label(person):
    def label(s):
        if not s:
            return ''
        return s if len(s)==4 else ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'][int(s[5:])-1]+' '+s[:4]
    start=label(person.get('joined',''))
    return start+(' to Present' if person.get('current') is True else ' to '+label(person['left']) if person.get('current') is False else '') if start else 'Not available'


class SourcerCandidates:
    MAX_ROWS = 5000
    MAX_BYTES = 8 * 1024 * 1024

    def __init__(self, *, load, save):
        self.load, self.save = load, save
        self.lock = threading.RLock()

    def _read(self):
        try:
            value = self.load()
            if not value:
                return {}
            rows = value['rows']
            if value['version'] != 1 or not isinstance(rows, dict) or len(rows) > self.MAX_ROWS or len(json.dumps({'version':1,'rows':rows}).encode()) > self.MAX_BYTES:
                raise ValueError()
            for key, row in rows.items():
                if not re.fullmatch(r'[a-f0-9]{64}', key) or not isinstance(row, dict):
                    raise ValueError()
                for field, limit in {'name':150,'title':500,'location':240,'url':400,'brief':4000,'run_id':32,'saved_at':40,'notes':2000,'read':20}.items():
                    if not isinstance(row.get(field), str) or len(row[field]) > limit or '\x00' in row[field]:
                        raise ValueError()
                if type(row.get('revision')) is not int or row['revision'] < 1 or type(row.get('score')) is not int or not 0 <= row['score'] <= 100:
                    raise ValueError()
                if row['read'] not in {'Search card', 'Profile detail'}:
                    raise ValueError()
                url=urlsplit(row['url'])
                if (url.scheme!='https' or url.netloc!='www.linkedin.com' or url.query or url.fragment or
                        not re.fullmatch(r'/(?:in/[^/]+|talent/(?:hire/)?profile/[^/]+)',url.path) or
                        re.search(r'[\x00-\x20\\]',unquote(url.path)) or not re.fullmatch(r'[a-f0-9]{32}', row['run_id'])):
                    raise ValueError()
                if key != sha256((row['run_id']+'|'+row['url']).encode()).hexdigest():
                    raise ValueError()
                for field, maximum in [('evidence',4), ('unknowns',6)]:
                    if not isinstance(row.get(field),list) or len(row[field]) > maximum or any(not isinstance(s,str) or len(s)>250 for s in row[field]):
                        raise ValueError()
            return deepcopy(rows)
        except Exception:
            raise PeopleError('Saved candidates could not be read securely. Keep the existing files and check local storage.') from None

    def _write(self, rows):
        value = {'version':1,'rows':rows}
        if len(rows) > self.MAX_ROWS or len(json.dumps(value).encode()) > self.MAX_BYTES:
            raise PeopleError('The candidate dashboard is full. Export the current search; its results are still shown in Search.')
        try:
            self.save(value)
        except Exception:
            raise PeopleError('Candidates could not be saved to the dashboard. Export the current search before closing it.') from None

    def remember(self, run, candidates, stage):
        with self.lock:
            rows = self._read()
            for person in candidates:
                if '\x00' in person['name'] or '\x00' in run.brief:
                    raise PeopleError('This assessment contains unsupported text. Export the current search; existing saved candidates were kept.')
                key = sha256((run.id+'|'+person['url']).encode()).hexdigest()
                old = rows.get(key)
                if old and old['read'] == 'Profile detail' and stage != 'detail':
                    continue
                rows[key] = {key:person.get(key, '') for key in ['name','title','location','url','score','evidence','unknowns']}
                rows[key].update(run_id=run.id, brief=run.brief,
                    read='Profile detail' if stage == 'detail' else 'Search card',
                    saved_at=datetime.now(timezone.utc).isoformat(timespec='seconds'),
                    notes=old['notes'] if old else '', revision=old['revision']+1 if old else 1)
            if candidates:
                self._write(rows)

    def list(self):
        with self.lock:
            return sorted(({**row,'key':key} for key,row in self._read().items()),
                          key=lambda r:(r['saved_at'],r['run_id'],r['score']), reverse=True)

    def note(self, key, revision, notes):
        if not isinstance(key,str) or not re.fullmatch(r'[a-f0-9]{64}',key) or type(revision) is not int or not isinstance(notes,str) or len(notes)>2000 or '\x00' in notes:
            raise PeopleError('Choose a saved candidate and a note of up to 2000 characters.')
        with self.lock:
            rows = self._read()
            if key not in rows or rows[key]['revision'] != revision:
                raise PeopleError('This candidate changed. Copy your note, refresh the dashboard and try again.')
            rows[key].update(notes=notes,revision=revision+1)
            self._write(rows)
            return {**rows[key],'key':key}


def hr_people(rows, cached=None):
    """One current saved observation per person/company, with exact action source."""
    found = {}
    for row in sorted(rows, key=lambda r:(r.get('hr') or {}).get('scanned_at',''), reverse=True):
        hr = row.get('hr')
        if row.get('deleted') or not hr:
            continue
        for person in hr['people']:
            key = sha256((hr['company_id']+'|'+person['url'].rstrip('/').lower()).encode()).hexdigest()
            result=cached(row,person) if cached else None
            rank=lambda value:3 if value and value.get('state')=='found' else 2 if value and value.get('state') in {'pending','unconfirmed'} else 1 if value else 0
            if key in found:
                found[key]['ads'] += 1
                if rank(result)>rank(found[key]['email_result']):
                    found[key].update(email_result=result,email_id=row['id'],email_revision=row['revision'])
                continue
            found[key] = {**person, 'key':key, 'id':row['id'], 'revision':row['revision'],
                'company':row['company'], 'company_url':hr['company_url'], 'company_id':hr['company_id'],
                'country':hr['country'], 'scanned_at':hr['scanned_at'], 'ads':1,
                'job_url':row['url'], 'folder':row.get('folder',''), 'employment':employment_label(person),
                'email_result':result,'email_id':row['id'],'email_revision':row['revision']}
    return list(found.values())
